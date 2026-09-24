"""Exporter tests: the Excel file must really have the promised formatting."""

from __future__ import annotations

import csv
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from bookscraper.exporters import COLUMNS, export, resolve_outputs
from bookscraper.models import Book


def make_book(i: int, **kw) -> Book:
    data = dict(
        title=f"Book {i}",
        price=10.0 + i,
        currency="GBP",
        availability=i,
        in_stock=True,
        rating=(i % 5) + 1,
        upc=f"upc{i:04d}",
        category="Mystery" if i % 2 else "Poetry",
        description="Описание — «кириллица» + £ sign",
        image_url=f"https://books.toscrape.com/media/{i}.jpg",
        product_url=f"https://books.toscrape.com/catalogue/book_{i}/index.html",
        num_reviews=0,
        scraped_at="2026-01-01T00:00:00+00:00",
    )
    data.update(kw)
    return Book(**data)


@pytest.fixture
def books() -> list[Book]:
    return [make_book(i) for i in range(1, 6)] + [make_book(6, price=None, rating=None)]


def test_xlsx_formatting(tmp_path, books):
    path = export(books, tmp_path / "out.xlsx", {"categories": "all", "limit": 6})
    wb = load_workbook(path)
    assert wb.sheetnames == ["Books", "Summary", "About"]
    ws = wb["Books"]

    headers = [c.value for c in ws[1]]
    assert headers == [c.header for c in COLUMNS]
    assert ws["A1"].font.bold and ws["A1"].fill.fgColor.rgb.endswith("1F3864")
    assert ws.freeze_panes == "B2"
    assert ws.auto_filter.ref == f"A1:K{len(books) + 1}"
    assert ws.max_row == len(books) + 1
    assert ws.column_dimensions["A"].width == 45

    col = {c.key: i for i, c in enumerate(COLUMNS, start=1)}
    url_cell = ws.cell(row=2, column=col["product_url"])
    assert url_cell.hyperlink.target == books[0].product_url
    assert ws.cell(row=2, column=col["title"]).hyperlink.target == books[0].product_url
    assert ws.cell(row=2, column=col["image_url"]).hyperlink.target == books[0].image_url
    price = ws.cell(row=2, column=col["price"])
    assert price.value == 11.0 and "£" in price.number_format
    assert ws.cell(row=7, column=col["price"]).value is None  # missing values stay empty
    assert ws.conditional_formatting  # data bar + colour scale


def test_xlsx_summary_and_about(tmp_path, books):
    wb = load_workbook(export(books, tmp_path / "out.xlsx", {"limit": 6}))
    rows = list(wb["Summary"].iter_rows(values_only=True))
    assert rows[0][0] == "Category"
    by_cat = {r[0]: r for r in rows[1:]}
    assert by_cat["Mystery"][1] == 3 and by_cat["Poetry"][1] == 3
    assert by_cat["TOTAL"][1] == 6
    assert by_cat["TOTAL"][5] == sum(b.availability for b in books)
    about = {r[0]: r[1] for r in wb["About"].iter_rows(values_only=True) if r[0]}
    assert "Демо-проект / Demo project" in wb["About"]["A2"].value
    assert about["Rows"] == 6 and about["Limit"] == 6


def test_xlsx_empty_list(tmp_path):
    wb = load_workbook(export([], tmp_path / "empty.xlsx"))
    assert wb["Books"].max_row == 1


def test_csv_utf8_bom_and_roundtrip(tmp_path, books):
    path = export(books, tmp_path / "out.csv")
    raw = path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")  # BOM for Excel
    with path.open(encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    assert len(rows) == len(books)
    assert rows[0]["Title"] == "Book 1"
    assert rows[0]["Description"] == "Описание — «кириллица» + £ sign"
    assert rows[-1]["Price"] == ""  # None -> empty cell


def test_json_roundtrip(tmp_path, books):
    path = export(books, tmp_path / "out.json")
    data = json.loads(path.read_text(encoding="utf-8"))
    assert len(data) == len(books)
    assert data[0] == books[0].to_dict()
    assert "кириллица" in path.read_text(encoding="utf-8")  # ensure_ascii=False


def test_resolve_outputs():
    assert resolve_outputs(["out/books"]) == [Path("out/books.xlsx"), Path("out/books.csv"),
                                              Path("out/books.json")]
    assert resolve_outputs(["a.csv", "a.CSV", "b.json"]) == [Path("a.csv"), Path("a.CSV"), Path("b.json")]
    assert resolve_outputs(["x.xlsx", "x.xlsx"]) == [Path("x.xlsx")]


def test_xlsx_keeps_formula_like_text_as_text(tmp_path):
    book = make_book(1, title="=HYPERLINK(\"http://evil.test\")", description="bad\x01chars\x1f here")
    wb = load_workbook(export([book], tmp_path / "t.xlsx"))
    ws = wb["Books"]
    col = {c.key: i for i, c in enumerate(COLUMNS, start=1)}
    title = ws.cell(row=2, column=col["title"])
    assert title.data_type == "s" and title.value.startswith("=HYPERLINK")
    assert ws.cell(row=2, column=col["description"]).value == "badchars here"


def test_csv_semicolon_delimiter(tmp_path, books):
    path = export(books, tmp_path / "out.csv", csv_delimiter=";")
    with path.open(encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh, delimiter=";"))
    assert rows[0]["Title"] == "Book 1" and rows[0]["Category"] == "Mystery"
    assert rows[0]["Price"] == "11,0"  # decimal comma, as a Russian-locale Excel expects
    assert rows[0]["In stock (qty)"] == "1"
    assert path.read_text(encoding="utf-8-sig").splitlines()[0].startswith("Title;Category;Price")


def test_resolve_outputs_directory(tmp_path):
    assert resolve_outputs([str(tmp_path)]) == [tmp_path / "books.xlsx", tmp_path / "books.csv",
                                                tmp_path / "books.json"]
