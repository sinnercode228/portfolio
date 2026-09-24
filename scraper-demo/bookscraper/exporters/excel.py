"""Client-ready Excel report (openpyxl).

Sheet "Books":   formatted header, autofilter, frozen header + title column,
                 column widths, number formats, clickable hyperlinks,
                 data bars on stock, colour scale on rating.
Sheet "Summary": per-category stats.
Sheet "About":   source, parameters, timestamp, demo note.
"""

from __future__ import annotations

from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from statistics import mean

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from bookscraper.exporters.tabular import COLUMNS, row_values
from bookscraper.models import Book

HEADER_FILL = PatternFill("solid", start_color="1F3864", end_color="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_BORDER = Border(bottom=Side(style="medium", color="0B1A33"))
BAND_FILL = PatternFill("solid", start_color="F2F5FA", end_color="F2F5FA")
LINK_FONT = Font(color="0563C1", underline="single")
DEMO_NOTE = "Демо-проект / Demo project — data from the books.toscrape.com scraping sandbox."
EXCEL_CELL_LIMIT = 32767  # max characters Excel accepts in one cell


def _safe(value: object) -> object:
    """Make scraped text safe for a cell: drop control characters openpyxl rejects,
    respect Excel's per-cell length limit. (Leading '=' is handled by the caller.)"""
    if not isinstance(value, str):
        return value
    value = ILLEGAL_CHARACTERS_RE.sub("", value)
    if len(value) > EXCEL_CELL_LIMIT:
        value = value[: EXCEL_CELL_LIMIT - 1] + "…"
    return value


def _set_text_value(ws: Worksheet, row: int, column: int, value: object):
    """Write a value; strings are always stored as text, never as formulas
    (a scraped title like "=SUM(A1)" must stay a title)."""
    cell = ws.cell(row=row, column=column, value=_safe(value))
    if isinstance(cell.value, str) and cell.data_type == "f":
        cell.data_type = "s"
    return cell


def _style_header(ws: Worksheet, n_cols: int) -> None:
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = HEADER_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30


def _books_sheet(ws: Worksheet, books: list[Book]) -> None:
    ws.title = "Books"
    ws.append([c.header for c in COLUMNS])
    _style_header(ws, len(COLUMNS))

    title_idx = next(i for i, c in enumerate(COLUMNS) if c.key == "title")
    url_idx = next(i for i, c in enumerate(COLUMNS) if c.key == "product_url")

    for r, book in enumerate(books, start=2):
        values = row_values(book)
        for c, (col, value) in enumerate(zip(COLUMNS, values, strict=True), start=1):
            cell = _set_text_value(ws, r, c, value)
            cell.alignment = Alignment(vertical="top")
            if col.number_format and isinstance(value, (int, float)):
                cell.number_format = col.number_format
            if col.hyperlink and value:
                cell.hyperlink = str(value)
                cell.font = LINK_FONT
            if r % 2 == 0:
                cell.fill = BAND_FILL
        # the title itself links to the product page too
        product_url = values[url_idx]
        if product_url:
            title_cell = ws.cell(row=r, column=title_idx + 1)
            title_cell.hyperlink = str(product_url)
            title_cell.font = LINK_FONT

    for i, col in enumerate(COLUMNS, start=1):
        width = col.width
        if col.hyperlink:  # make URL columns wide enough for the longest URL (capped)
            longest = max((len(str(getattr(b, col.key) or "")) for b in books), default=0)
            width = max(width, min(longest + 2, 42))
        ws.column_dimensions[get_column_letter(i)].width = width

    last_row = max(2, len(books) + 1)
    last_col = get_column_letter(len(COLUMNS))
    ws.auto_filter.ref = f"A1:{last_col}{last_row}"
    ws.freeze_panes = "B2"  # header row + Title column stay visible

    if books:
        for key, rule in (
            ("availability", DataBarRule(start_type="min", end_type="max", color="5B9BD5")),
            ("rating", ColorScaleRule(start_type="num", start_value=1, start_color="F8696B",
                                      mid_type="num", mid_value=3, mid_color="FFEB84",
                                      end_type="num", end_value=5, end_color="63BE7B")),
        ):
            idx = next(i for i, c in enumerate(COLUMNS, start=1) if c.key == key)
            letter = get_column_letter(idx)
            ws.conditional_formatting.add(f"{letter}2:{letter}{last_row}", rule)

    ws.print_title_rows = "1:1"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _summary_sheet(ws: Worksheet, books: list[Book]) -> None:
    headers = ["Category", "Books", "Avg price", "Min price", "Max price", "Units in stock", "Avg rating"]
    ws.append(headers)
    _style_header(ws, len(headers))

    groups: dict[str, list[Book]] = defaultdict(list)
    for book in books:
        groups[book.category or "(none)"].append(book)

    def stats(items: list[Book]) -> list[object]:
        prices = [b.price for b in items if b.price is not None]
        ratings = [b.rating for b in items if b.rating is not None]
        return [
            len(items),
            round(mean(prices), 2) if prices else None,
            min(prices) if prices else None,
            max(prices) if prices else None,
            sum(b.availability for b in items),
            round(mean(ratings), 2) if ratings else None,
        ]

    for name in sorted(groups, key=lambda n: (-len(groups[n]), n)):
        row = ws.max_row + 1
        for col, value in enumerate([name, *stats(groups[name])], start=1):
            _set_text_value(ws, row, col, value)
    ws.append(["TOTAL", *stats(books)])
    total_row = ws.max_row
    for cell in ws[total_row]:
        cell.font = Font(bold=True)
        cell.border = Border(top=Side(style="thin"))

    formats = [None, "0", '"£"#,##0.00', '"£"#,##0.00', '"£"#,##0.00', "#,##0", "0.00"]
    widths = [24, 9, 12, 12, 12, 15, 12]
    for i, (fmt, width) in enumerate(zip(formats, widths, strict=True), start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
        if fmt:
            for row in range(2, total_row + 1):
                ws.cell(row=row, column=i).number_format = fmt
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(2, total_row - 1)}"


def _about_sheet(ws: Worksheet, books: list[Book], meta: dict[str, object]) -> None:
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 90
    ws["A1"] = "Books scraper — demo report"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = DEMO_NOTE
    ws["A2"].font = Font(italic=True, color="7F7F7F")
    rows: list[tuple[str, object]] = [
        ("Source", meta.get("source", "https://books.toscrape.com/")),
        ("Generated (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")),
        ("Rows", len(books)),
    ]
    rows += [(str(k).replace("_", " ").capitalize(), v) for k, v in meta.items() if k != "source"]
    for i, (key, value) in enumerate(rows, start=4):
        ws.cell(row=i, column=1, value=key).font = Font(bold=True)
        _set_text_value(ws, i, 2, str(value) if isinstance(value, (list, tuple, dict)) else value)
        ws.cell(row=i, column=2).alignment = Alignment(horizontal="left")


def write_xlsx(books: list[Book], path: Path, *, meta: dict[str, object] | None = None) -> Path:
    wb = Workbook()
    _books_sheet(wb.active, books)
    _summary_sheet(wb.create_sheet("Summary"), books)
    _about_sheet(wb.create_sheet("About"), books, dict(meta or {}))
    wb.properties.title = "Books scraper — demo report"
    wb.properties.creator = "bookscraper-demo"
    wb.properties.description = DEMO_NOTE
    wb.active = 0
    wb.save(path)
    return path
