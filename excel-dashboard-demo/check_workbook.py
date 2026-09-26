#!/usr/bin/env python3
"""Verify generated workbooks without Excel.

For every file:
  1. lint: allowed Excel-2013 functions only, brackets/quotes, sheets, tables, columns, names;
  2. evaluate every formula with pycel and count Excel errors (#VALUE!, #REF!, #NAME?, ...);
  3. if LibreOffice (soffice) is installed, also recalculate headlessly and scan for errors.

Example:
    python check_workbook.py output/dashboard_sample.xlsx output/invoice_sample.xlsx
"""

from __future__ import annotations

import argparse
import glob
import sys
from pathlib import Path

from openpyxl import load_workbook

from xldash.formula_check import (ERROR_VALUES, evaluate_workbook, find_errors, lint_workbook,
                                  recalc_with_libreoffice)


def check(path: str) -> bool:
    print(f"== {path}")
    lint = lint_workbook(path)
    print(f"   formulas: {lint.formulas}")
    print(f"   functions: {', '.join(f'{k}×{v}' for k, v in sorted(lint.functions.items()))}")
    for w in lint.warnings:
        print(f"   warning: {w}")
    for p in lint.problems[:25]:
        print(f"   PROBLEM: {p}")
    values = evaluate_workbook(path)
    errors = find_errors(values)
    print(f"   evaluated with pycel: {len(values)} formulas, {len(errors)} error value(s)")
    for addr, err in list(errors.items())[:25]:
        print(f"   ERROR: {addr} = {err}")
    ok = lint.ok and not errors

    recalculated = recalc_with_libreoffice(path)
    if recalculated is None:
        print("   LibreOffice: not installed or conversion failed - skipped")
    else:
        wb = load_workbook(recalculated, data_only=True)
        lo_errors = [f"{ws.title}!{c.coordinate}" for ws in wb.worksheets for row in ws.iter_rows()
                     for c in row if isinstance(c.value, str) and c.value in ERROR_VALUES]
        print(f"   LibreOffice recalculation: {len(lo_errors)} error value(s)")
        ok = ok and not lo_errors
    print("   RESULT:", "OK" if ok else "FAILED")
    return ok


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("files", nargs="+", help="workbooks; wildcards like output/*.xlsx also work on Windows")
    args = p.parse_args(argv)
    files: list[str] = []
    for pattern in args.files:            # cmd.exe does not expand wildcards itself
        matches = sorted(glob.glob(pattern)) if glob.has_magic(pattern) else [pattern]
        files += matches
        if not matches or not all(Path(f).is_file() for f in matches):
            print(f"error: no such workbook: {pattern}", file=sys.stderr)
            return 2
    results = [check(f) for f in files]
    return 0 if all(results) else 1


if __name__ == "__main__":
    sys.exit(main())
