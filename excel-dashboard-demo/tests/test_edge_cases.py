"""Regression tests for inputs a real client file can contain (found in review)."""

import importlib
from datetime import date

import pytest
from openpyxl import load_workbook

from conftest import PERIOD, find_row
from xldash.dashboard import LABELS, build_dashboard
from xldash.formula_check import evaluate_workbook, find_errors, lint_workbook
from xldash.invoice import L10N
from xldash.production import load_production_csv
from xldash.worklog import (Rate, WorkEntry, invoice_reference, load_rates_csv,
                            match_rate_names)

HEADER = "date,shift,department,machine,employee,run_hours,halt_minutes,halt_reason\n"


def _dashboard(tmp_path, body: str, lang: str = "en"):
    src = tmp_path / "log.csv"
    src.write_text(HEADER + body, encoding="utf-8")
    records, report = load_production_csv(src, lang=lang)
    out = build_dashboard(records, tmp_path / "d.xlsx", lang=lang, report=report)
    return out, evaluate_workbook(out)


def test_log_without_any_halt_still_builds(tmp_path):
    out, values = _dashboard(tmp_path, "2026-09-01,Day,Assembly,A1,Ann,8,0,\n"
                                       "2026-09-02,Day,Assembly,A1,Ann,7.5,0,\n")
    assert lint_workbook(out).ok
    assert find_errors(values) == {}
    assert values["Dashboard!B6"] == pytest.approx(15.5)
    assert values["Dashboard!F6"] == pytest.approx(1.0)
    assert values["Dashboard!P6"] == "–"                  # no top reason


def test_single_halt_reason_is_ranked(tmp_path):
    # one reason -> one-cell ranges such as $C$9:$C$9 in MATCH/INDEX
    _, values = _dashboard(tmp_path, "2026-09-01,Day,Assembly,A1,Ann,7.5,30,Jam\n")
    assert find_errors(values) == {}
    assert values["Dashboard!P6"] == "Jam"


def test_russian_import_log_is_in_russian(tmp_path):
    src = tmp_path / "ru.csv"
    src.write_text("Дата;Смена;Цех;Оборудование;Сотрудник;Часы работы;Простой, мин;Причина простоя\n"
                   "01.09.2026;Ночная;Сварка;WLD-01;Иванов И.;7,5;30;Поломка\n"
                   "01.09.2026;Ночная;Сварка;WLD-01;Иванов И.;7,5;30;Поломка\n"
                   "45.13.2026;Ночная;Сварка;WLD-01;Иванов И.;8;0;\n", encoding="utf-8")
    records, report = load_production_csv(src, lang="ru")
    out = build_dashboard(records, tmp_path / "ru.xlsx", lang="ru", report=report)
    ws = load_workbook(out)[LABELS["ru"]["sheets"]["log"]]
    texts = [str(v) for row in ws.iter_rows(values_only=True) for v in row if v is not None]
    joined = " | ".join(texts)
    assert "Принято строк: 1 из 3" in joined
    for english in ("rows accepted", "duplicate", "unrecognised", "rejected_", "dropped_"):
        assert english not in joined, english
    assert ws.auto_filter.ref


def test_blank_increment_means_no_rounding(invoice, worklog, lang):
    path, wb = invoice
    entries, rates, _ = worklog
    T = L10N[lang]
    rows = {name: 5 + i for i, (name, *_rest) in enumerate(T["set_rows"])}
    for inc in (0, None):                                 # user typed 0 or cleared the cell
        values = evaluate_workbook(path, overrides={f"{T['sheets']['settings']}!C{rows['billInc']}": inc})
        assert find_errors(values) == {}
        ref = invoice_reference(entries, rates, *PERIOD, increment=0)
        inv = wb[T["sheets"]["invoice"]]
        sub = find_row(inv, 5, T["subtotal"])
        assert values[f"{T['sheets']['invoice']}!G{sub}"] == pytest.approx(ref["subtotal"])


def test_rate_names_are_matched_case_insensitively():
    rates = [Rate("Mira Solberg", "Lead", 55, 30)]
    entries = [WorkEntry(date(2026, 9, 1), "mira solberg", "P", "t", 1.0, True),
               WorkEntry(date(2026, 9, 1), "Somebody Else", "P", "t", 1.0, True)]
    fixed = match_rate_names(entries, rates)
    assert [e.employee for e in fixed] == ["Mira Solberg", "Somebody Else"]


@pytest.mark.parametrize("content, message", [
    ("", "file is empty"),
    ("employee,role,bill_rate,pay_rate\n", "no rates found"),
    ("employee,role,bill_rate,pay_rate\nMira,Lead,55\n", "line 2: bad rate"),
    ("employee,role,bill_rate,pay_rate\nMira,Lead,-5,10\n", "negative rate"),
    ("employee,role,bill_rate\nMira,Lead,55\n", "missing required column"),
])
def test_bad_rates_file_is_a_clear_error(tmp_path, content, message):
    path = tmp_path / "rates.csv"
    path.write_text(content, encoding="utf-8")
    with pytest.raises(ValueError, match=message):
        load_rates_csv(path)


def _main(module):
    return importlib.import_module(module).main


@pytest.mark.parametrize("extra", [["--from", "2026-13-01"], ["--increment", "-1"], ["--tax", "20"],
                                   ["--invoice-date", "tomorrow"]])
def test_invoice_cli_rejects_bad_arguments(tmp_path, extra):
    _main("generate_sample_data")(["--out-dir", str(tmp_path), "--lang", "en"])
    with pytest.raises(SystemExit) as exc:
        _main("build_invoice")([str(tmp_path / "work_log_sample.csv"), str(tmp_path / "rates_sample.csv"),
                                "-o", str(tmp_path / "x.xlsx"), *extra])
    assert exc.value.code == 2


def test_invoice_cli_accepts_russian_dates_and_zero_increment(tmp_path, capsys):
    _main("generate_sample_data")(["--out-dir", str(tmp_path), "--lang", "en"])
    out = tmp_path / "x.xlsx"
    assert _main("build_invoice")([str(tmp_path / "work_log_sample.csv"), str(tmp_path / "rates_sample.csv"),
                                   "-o", str(out), "--from", "01.09.2026", "--to", "15.09.2026",
                                   "--increment", "0"]) == 0
    assert "Period 2026-09-01 .. 2026-09-15" in capsys.readouterr().out
    assert find_errors(evaluate_workbook(out)) == {}


def test_cli_reports_unwritable_output(tmp_path, capsys):
    _main("generate_sample_data")(["--out-dir", str(tmp_path), "--lang", "en"])
    assert _main("build_dashboard")([str(tmp_path / "production_log_sample.csv"), "-o", str(tmp_path)]) == 2
    assert "cannot write" in capsys.readouterr().err


def test_generator_work_log_follows_start_and_days(tmp_path):
    _main("generate_sample_data")(["--out-dir", str(tmp_path), "--lang", "en",
                                   "--start", "2026-01-05", "--days", "10"])
    lines = (tmp_path / "work_log_sample.csv").read_text(encoding="utf-8-sig").splitlines()[1:]
    dates = sorted(line.split(",")[0] for line in lines)
    assert "2025-12-29" <= dates[0] and dates[-1] <= "2026-01-17"


def test_invoice_project_table_shows_average_rate(invoice, lang):
    path, wb = invoice
    T = L10N[lang]
    i = T["sheets"]["invoice"]
    inv = wb[i]
    values = evaluate_workbook(path, sheets=[i])
    first = find_row(inv, 2, T["by_project"]) + 2         # section title, header, first project
    total = find_row(inv, 3, T["total_row"], first)
    for r in range(first, total + 1):
        hours, rate, amount = (values[f"{i}!{c}{r}"] for c in "EFG")
        assert hours > 0 and rate == pytest.approx(amount / hours)


def test_cached_values_for_previews(tmp_path, capsys):
    """Viewers that do not calculate (Quick Look, phone previews) read the cached <v> values."""
    _main("generate_sample_data")(["--out-dir", str(tmp_path), "--lang", "en"])
    out = tmp_path / "inv.xlsx"
    assert _main("build_invoice")([str(tmp_path / "work_log_sample.csv"), str(tmp_path / "rates_sample.csv"),
                                   "-o", str(out)]) == 0
    assert "Stored computed values of" in capsys.readouterr().out
    formulas = load_workbook(out)["Invoice"]
    cached = load_workbook(out, data_only=True)["Invoice"]
    assert formulas["B2"].value == "=ctrName"                     # formulas are kept
    assert cached["B2"].value.startswith("Brindlecote Digital")
    values = evaluate_workbook(out)
    assert find_errors(values) == {}
    sub = find_row(formulas, 5, "Subtotal")
    assert cached[f"G{sub}"].value == pytest.approx(values[f"Invoice!G{sub}"]) and cached[f"G{sub}"].value > 0
    assert lint_workbook(out).ok

    raw = tmp_path / "raw.xlsx"
    _main("build_dashboard")([str(tmp_path / "production_log_sample.csv"), "-o", str(raw),
                              "--no-preview-values"])
    assert load_workbook(raw, data_only=True)["Dashboard"]["B6"].value is None
