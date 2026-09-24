"""Shared fixtures: sample CSVs and workbooks are generated once per test session."""

from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from xldash.dashboard import build_dashboard  # noqa: E402
from xldash.invoice import build_invoice  # noqa: E402
from xldash.production import load_production_csv  # noqa: E402
from xldash.sample_data import write_production_csv, write_rates_csv, write_worklog_csv  # noqa: E402
from xldash.worklog import load_rates_csv, load_worklog_csv  # noqa: E402

PERIOD = (date(2026, 9, 1), date(2026, 9, 30))


@pytest.fixture(scope="session")
def tmp_session(tmp_path_factory) -> Path:
    return tmp_path_factory.mktemp("xldash")


@pytest.fixture(scope="session", params=["en", "ru"])
def lang(request) -> str:
    return request.param


@pytest.fixture(scope="session")
def production(tmp_session, lang):
    """(records, report, csv_path) loaded from a messy sample log."""
    path = tmp_session / f"prod_{lang}.csv"
    write_production_csv(path, lang=lang, messy=True)
    records, report = load_production_csv(path, lang=lang)
    return records, report, path


@pytest.fixture(scope="session")
def dashboard(tmp_session, production, lang):
    records, report, _ = production
    out = build_dashboard(records, tmp_session / f"dashboard_{lang}.xlsx", lang=lang, report=report)
    return out, load_workbook(out)


@pytest.fixture(scope="session")
def worklog(tmp_session, lang):
    log = tmp_session / f"log_{lang}.csv"
    rates = tmp_session / f"rates_{lang}.csv"
    write_worklog_csv(log, lang=lang)
    write_rates_csv(rates, lang=lang)
    entries, report = load_worklog_csv(log)
    return entries, load_rates_csv(rates), report


@pytest.fixture(scope="session")
def invoice(tmp_session, worklog, lang):
    entries, rates, _ = worklog
    out = build_invoice(entries, rates, tmp_session / f"invoice_{lang}.xlsx", lang=lang,
                        period_from=PERIOD[0], period_to=PERIOD[1], increment=0.25, tax_rate=0.0)
    return out, load_workbook(out)


def find_row(ws, col: int, value, start: int = 1) -> int:
    """First row >= start whose cell in `col` equals `value`."""
    for r in range(start, ws.max_row + 1):
        if ws.cell(row=r, column=col).value == value:
            return r
    raise LookupError(f"{value!r} not found in column {col} of {ws.title}")
