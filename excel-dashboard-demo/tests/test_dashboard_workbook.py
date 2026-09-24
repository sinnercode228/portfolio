"""The dashboard workbook: structure, formula lint and real formula evaluation (pycel).

The evaluation tests change the Settings inputs (like a user picking from the
dropdowns) and compare every KPI / table value with the Python reference.
"""

from datetime import date, datetime

import pytest

from conftest import find_row
from xldash.dashboard import LABELS, TABLE
from xldash.formula_check import evaluate_workbook, find_errors, lint_workbook
from xldash.production import dimension_values, summarize


def _sheets(lang):
    return LABELS[lang]["sheets"]


def test_structure(dashboard, production, lang):
    path, wb = dashboard
    records, _, _ = production
    S = _sheets(lang)
    assert wb.sheetnames == [S["dash"], S["halts"], S["settings"], S["data"], S["log"], S["lists"]]
    assert wb[S["lists"]].sheet_state == "hidden"

    data = wb[S["data"]]
    table = data.tables[TABLE]
    assert table.ref == f"A1:H{len(records) + 1}"
    assert [c.value for c in data[1]] == LABELS[lang]["columns"]
    assert isinstance(data["A2"].value, datetime)

    settings = wb[S["settings"]]
    sources = sorted(dv.formula1 for dv in settings.data_validations.dataValidation if dv.type == "list")
    assert sources == ["lstDates", "lstDates", "lstDept", "lstShift"]
    for name in ("selShift", "selDept", "selFrom", "selTo", "tgtAvail", "critShift", "critDept",
                 "dFrom", "dTo", "lstShift", "lstDept", "lstDates"):
        assert name in wb.defined_names, name

    assert len(wb[S["dash"]]._charts) == 3
    assert len(wb[S["halts"]]._charts) == 2
    assert wb.calculation.fullCalcOnLoad


def test_lint_is_clean(dashboard):
    result = lint_workbook(dashboard[0])
    assert result.ok, result.problems[:10]
    assert result.formulas > 500
    assert {"SUMIFS", "COUNTIFS", "INDEX", "MATCH", "LARGE"} <= set(result.functions)
    assert not {"XLOOKUP", "FILTER", "UNIQUE", "SORT", "LET"} & set(result.functions)


@pytest.fixture(scope="module")
def evaluate(dashboard, lang):
    cache = {}

    def run(**settings):
        S = _sheets(lang)
        cells = {"shift": "C5", "dept": "C6", "date_from": "C7", "date_to": "C8"}
        overrides = {f"{S['settings']}!{cells[k]}": v for k, v in settings.items()}
        key = tuple(sorted((k, str(v)) for k, v in overrides.items()))
        if key not in cache:
            cache[key] = evaluate_workbook(dashboard[0], overrides=overrides,
                                           sheets=[S["dash"], S["halts"], S["settings"], S["lists"]])
        return cache[key]
    return run


def _check_kpis(values, ref, S):
    d = S["dash"]
    assert values[f"{d}!B6"] == pytest.approx(ref["run_hours"])
    assert values[f"{d}!D6"] == pytest.approx(ref["halt_hours"])
    if ref["availability"] is None:
        assert values[f"{d}!F6"] == ""
    else:
        assert values[f"{d}!F6"] == pytest.approx(ref["availability"])
    assert values[f"{d}!J6"] == ref["rows"]
    assert values[f"{d}!M6"] == ref["halt_events"]


def test_no_formula_errors_default_filters(evaluate):
    values = evaluate()
    assert find_errors(values) == {}


def _scenarios(records, lang):
    dims = dimension_values(records)
    all_ = LABELS[lang]["all"]
    return [
        ({}, {}),
        ({"shift": dims["shift"][-1]}, {"shift": dims["shift"][-1]}),
        ({"dept": dims["department"][-1]}, {"department": dims["department"][-1]}),
        ({"shift": dims["shift"][0], "dept": dims["department"][0], "date_from": datetime(2026, 9, 7),
          "date_to": datetime(2026, 9, 13)},
         {"shift": dims["shift"][0], "department": dims["department"][0],
          "date_from": date(2026, 9, 7), "date_to": date(2026, 9, 13)}),
        ({"shift": all_, "dept": all_, "date_from": datetime(2026, 9, 20)}, {"date_from": date(2026, 9, 20)}),
        ({"shift": ""}, {}),                                           # blank input = All
    ]


@pytest.mark.parametrize("scenario", range(6))
def test_kpis_follow_the_settings_filters(evaluate, production, lang, scenario):
    records, _, _ = production
    settings, filters = _scenarios(records, lang)[scenario]
    values = evaluate(**settings)
    assert find_errors(values) == {}
    _check_kpis(values, summarize(records, **filters), _sheets(lang))


def test_empty_period_gives_zeros_not_errors(evaluate, production, lang):
    records, _, _ = production
    S = _sheets(lang)
    values = evaluate(date_from=datetime(2026, 9, 20), date_to=datetime(2026, 9, 10))
    assert find_errors(values) == {}
    _check_kpis(values, summarize(records, date_from=date(2026, 9, 20), date_to=date(2026, 9, 10)), S)
    assert values[f"{S['settings']}!B11"].startswith("⚠")


def test_summary_tables_match_reference(evaluate, dashboard, production, lang):
    records, _, _ = production
    _, wb = dashboard
    S, L = _sheets(lang), LABELS[lang]
    dept = dimension_values(records)["department"][0]
    values = evaluate(dept=dept)
    ref = summarize(records, department=dept)
    ws = wb[S["dash"]]
    for key, group in (("by_machine", "by_machine"), ("by_employee", "by_employee"),
                       ("by_department", "by_department")):
        start = find_row(ws, 2, L[key])
        for name in dimension_values(records)[{"by_machine": "machine", "by_employee": "employee",
                                              "by_department": "department"}[key]]:
            r = find_row(ws, 2, name, start)
            g = ref[group].get(name, {"run": 0, "halt": 0, "shifts": 0, "availability": None})
            assert values[f"{S['dash']}!C{r}"] == pytest.approx(g["run"]), (key, name)
            assert values[f"{S['dash']}!D{r}"] == pytest.approx(g["halt"]), (key, name)
            assert values[f"{S['dash']}!F{r}"] == g["shifts"], (key, name)
            if g["availability"] is None:
                assert values[f"{S['dash']}!E{r}"] == ""
            else:
                assert values[f"{S['dash']}!E{r}"] == pytest.approx(g["availability"])
        total = find_row(ws, 2, L["total"], start)
        assert values[f"{S['dash']}!C{total}"] == pytest.approx(ref["run_hours"])


@pytest.mark.parametrize("settings", [{}, {"shift_index": -1}])
def test_halts_top_reasons_and_matrix(evaluate, dashboard, production, lang, settings):
    records, _, _ = production
    _, wb = dashboard
    S, L = _sheets(lang), LABELS[lang]
    kwargs, filters = {}, {}
    if settings:
        shift = dimension_values(records)["shift"][settings["shift_index"]]
        kwargs, filters = {"shift": shift}, {"shift": shift}
    values = evaluate(**kwargs)
    ref = summarize(records, **filters)
    ws = wb[S["halts"]]
    h = S["halts"]

    # top-N ranking equals the Python ranking
    top_start = find_row(ws, 2, L["h_top"]) + 2
    expected = sorted(ref["by_reason"].items(), key=lambda kv: -kv[1]["halt"])[:5]
    for k, (reason, g) in enumerate(expected):
        r = top_start + k
        assert values[f"{h}!C{r}"] == reason
        assert values[f"{h}!D{r}"] == pytest.approx(g["halt"])
        assert values[f"{h}!E{r}"] == g["events"]
    assert values[f"{S['dash']}!P6"] == expected[0][0]      # KPI tile "top reason"

    # halt hours per machine + top reason per machine (from the heat map)
    m_start = find_row(ws, 2, L["h_by_machine"])
    for machine, g in ref["by_machine"].items():
        r = find_row(ws, 2, machine, m_start)
        assert values[f"{h}!D{r}"] == pytest.approx(g["halt"])
        assert values[f"{h}!E{r}"] == g["events"]
        cells = {reason: v for (m, reason), v in ref["machine_reason"].items() if m == machine}
        if cells:
            assert values[f"{h}!H{r}"] == max(cells, key=cells.get)

    # matrix grand total = all halt hours
    mx = find_row(ws, 2, L["h_matrix"]) + 1
    total_row = find_row(ws, 2, L["total"], mx)
    last_col = 3 + len(dimension_values(records)["halt_reason"])
    from openpyxl.utils import get_column_letter
    assert values[f"{h}!{get_column_letter(last_col)}{total_row}"] == pytest.approx(ref["halt_hours"])


def test_date_dropdown_list_is_live(evaluate, production, lang):
    records, _, _ = production
    values = evaluate()
    lists = _sheets(lang)["lists"]
    days = sorted({r.date for r in records})
    n = (days[-1] - days[0]).days + 1
    assert values[f"{lists}!C2"] == (days[0] - date(1899, 12, 30)).days
    assert values[f"{lists}!C{n + 1}"] == (days[-1] - date(1899, 12, 30)).days
    assert values[f"{lists}!C{n + 2}"] == ""
