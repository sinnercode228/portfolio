"""Shared look: fonts, colours, number formats and small cell helpers."""

from __future__ import annotations

from openpyxl.chart import BarChart
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Font as DrawingFont, Paragraph, ParagraphProperties, RegularTextRun
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

FONT = "Arial"

# palette (validated categorical slots 1-2 + neutral chrome)
BLUE = "2A78D6"        # run hours / primary series
ORANGE = "EB6834"      # halt hours
INK = "0B0B0B"
INK_2 = "52514E"
MUTED = "898781"
GRID = "E1E0D9"
BASE = "C3C2B7"
PANEL = "F4F3EF"
HEADER_FILL = "1F3B57"  # dark slate for table headers / title bar
GOOD = "006300"        # success text (paired with an icon)
BAD = "D03B3B"         # critical (paired with an icon)
INPUT_FILL = "FFF4C2"  # soft yellow = editable cell
INPUT_FONT = "0000FF"  # blue = hard-coded input (financial-model convention)

thin = Side(style="thin", color=GRID)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
BOTTOM = Border(bottom=Side(style="thin", color=BASE))
TOP_DOUBLE = Border(top=Side(style="thin", color=INK_2), bottom=Side(style="double", color=INK_2))


def font(size=10, bold=False, color=INK, italic=False) -> Font:
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", start_color=hex_color, end_color=hex_color)


def set_widths(ws: Worksheet, widths: dict[str, float]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def base_sheet(ws: Worksheet, gridlines: bool = False) -> None:
    ws.sheet_view.showGridLines = gridlines
    ws.sheet_view.zoomScale = 100


def title_bar(ws: Worksheet, text: str, first_col: int, last_col: int, row: int = 1,
              subtitle: str | None = None) -> None:
    for c in range(first_col, last_col + 1):
        ws.cell(row=row, column=c).fill = fill(HEADER_FILL)
    cell = ws.cell(row=row, column=first_col, value=text)
    cell.font = font(16, True, "FFFFFF")
    cell.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[row].height = 30
    if subtitle:
        sub = ws.cell(row=row + 1, column=first_col, value=subtitle)
        sub.font = font(9, color=INK_2, italic=True)


def header_row(ws: Worksheet, row: int, first_col: int, labels: list[str],
               wrap: bool = True) -> None:
    for i, label in enumerate(labels):
        c = ws.cell(row=row, column=first_col + i, value=label)
        c.font = font(9, True, "FFFFFF")
        c.fill = fill(HEADER_FILL)
        c.alignment = Alignment(horizontal="center" if i else "left", vertical="center",
                                wrap_text=wrap, indent=0 if i else 1)
        c.border = BORDER
    ws.row_dimensions[row].height = 30


def section_title(ws: Worksheet, row: int, col: int, text: str, note: str | None = None) -> None:
    c = ws.cell(row=row, column=col, value=text)
    c.font = font(12, True, HEADER_FILL)
    if note:
        n = ws.cell(row=row, column=col + 2, value=note)
        n.font = font(8, color=MUTED, italic=True)


def body_cell(cell, fmt: str | None = None, bold=False, align="right", color=INK, fill_hex=None):
    cell.font = font(10, bold, color)
    cell.border = BORDER
    cell.alignment = Alignment(horizontal=align, vertical="center", indent=1 if align == "left" else 0)
    if fmt:
        cell.number_format = fmt
    if fill_hex:
        cell.fill = fill(fill_hex)


def input_cell(cell, fmt: str | None = None, align="center"):
    cell.font = font(11, True, INPUT_FONT)
    cell.fill = fill(INPUT_FILL)
    cell.border = Border(left=Side(style="thin", color="C9A800"), right=Side(style="thin", color="C9A800"),
                         top=Side(style="thin", color="C9A800"), bottom=Side(style="thin", color="C9A800"))
    cell.alignment = Alignment(horizontal=align, vertical="center")
    if fmt:
        cell.number_format = fmt


def col(n: int) -> str:
    return get_column_letter(n)


def style_bar_chart(chart: BarChart, title: str, *, horizontal: bool, stacked: bool = False,
                    y_title: str | None = None, number_format: str = "#,##0",
                    colors: tuple[str, ...] = (BLUE, ORANGE), legend: bool = True,
                    width: float = 16.0, height: float = 7.5, pct_axis: bool = False) -> BarChart:
    """Consistent, quiet chart chrome: thin bars, hairline grid, legend at the bottom."""
    chart.type = "bar" if horizontal else "col"
    chart.grouping = "stacked" if stacked else "clustered"
    if stacked:
        chart.overlap = 100
    chart.gapWidth = 60
    chart.title = chart_title(title)
    chart.style = 2
    chart.roundedCorners = False
    chart.width = width
    chart.height = height
    if legend:
        chart.legend.position = "b"
        chart.legend.overlay = False
        chart.legend.txPr = text_props(9, INK_2)
    else:
        chart.legend = None
    # openpyxl >= 3.1 hides axes unless told otherwise
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.y_axis.number_format = number_format
    chart.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=GRID))
    chart.y_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
    chart.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill=BASE))
    if pct_axis:
        chart.y_axis.scaling.min = 0
        chart.y_axis.scaling.max = 1
    if horizontal:
        chart.x_axis.scaling.orientation = "maxMin"   # first table row on top
        chart.y_axis.crosses = "max"                  # keep value axis at the bottom
    if y_title:
        chart.y_axis.title = y_title
    chart.x_axis.txPr = text_props(9, INK_2)
    chart.y_axis.txPr = text_props(9, MUTED)
    for i, s in enumerate(chart.series):
        color = colors[i % len(colors)]
        s.graphicalProperties.solidFill = color
        s.graphicalProperties.line.solidFill = color
    return chart


def _char(size_pt: float, color: str, bold: bool = False) -> CharacterProperties:
    return CharacterProperties(sz=int(size_pt * 100), b=bold, solidFill=color,
                               latin=DrawingFont(typeface=FONT))


def text_props(size_pt: float, color: str, bold: bool = False) -> RichText:
    """txPr for axes / legends / data labels: Arial, given size and colour."""
    cp = _char(size_pt, color, bold)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def chart_title(text: str) -> Title:
    """Small bold title that sits above the plot (never on top of the bars)."""
    cp = _char(11, HEADER_FILL, True)
    para = Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=text)])
    return Title(tx=Text(rich=RichText(p=[para])), overlay=False)
